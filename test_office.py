"""実際のExcel変換を通して、全単位の単価補完と既存単価の保持を確認する。"""
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from office import office_int, office_int_multi


class OfficeUnitPriceTest(unittest.TestCase):
    def test_unit_prices_in_both_formats(self):
        cases = [
            ("工事A", 26, "式", None, 390000, 15000),
            ("工事B", 1, "式", None, 30000, 30000),
            ("既存単価", 26, "式", 12345, 390000, 12345),
            ("文字列数値", "２６．００", " 式　", "　", "３９０，０００", 15000),
            ("端数", 3, "式", None, 100, 100 / 3),
            ("ゼロ数量", 0, "式", None, 100, None),
            ("数量不明", None, "式", None, 100, None),
            ("金額なし", 2, "式", None, None, None),
            ("検体単価", 3, "検体", None, 120000, 40000),
            ("人工単価", 7, "人工", None, 224000, 32000),
            ("台単価", 2, "台", None, 50000, 25000),
            ("面積単価", 2.5, "㎡", None, 10000, 4000),
            ("少量単価", 0.5, "m", None, 1000, 2000),
            ("単位なし", 4, None, None, 10000, 2500),
            ("その他単位1個", 1, "個", None, 30000, 30000),
            ("その他単位の端数", 3, "台", None, 100, 100 / 3),
            ("その他単位の既存単価", 7, "人工", 31000, 224000, 31000),
            ("ゼロ単価保持", 3, "検体", 0, 120000, 0),
            ("その他単位の文字列数値", "３．００", "検体", "　", "１２０，０００", 40000),
            ("その他単位のゼロ数量", 0, "個", None, 100, None),
            ("その他単位の数量不明", None, "個", None, 100, None),
            ("その他単位の不正数量", "不明", "個", None, 100, None),
            ("その他単位の負数数量", -2, "個", None, 100, None),
            ("その他単位の金額なし", 2, "個", None, None, None),
            ("小計", 26, "式", None, 390000, None),
        ]
        for multi in (False, True):
            with self.subTest(multi=multi), tempfile.TemporaryDirectory() as directory:
                wb = Workbook()
                ws = wb.active
                ws.title = "明細"
                ws.cell(1, 3, "品名")
                for row, (name, qty, unit, price, amount, _) in enumerate(cases, 2):
                    ws.cell(row, 2, row - 1)
                    ws.cell(row, 3, name)
                    ws.cell(row, 10, qty)
                    ws.cell(row, 11, unit)
                    ws.cell(row, 12, price)
                    ws.cell(row, 13 if multi else 14, amount)
                source = str(Path(directory) / "input.xlsx")
                wb.save(source)
                wb.close()
                output = office_int_multi(source, "明細") if multi else office_int(source)
                result = load_workbook(output, data_only=True)
                try:
                    actual = {row[2]: row[7] for row in result.active.iter_rows(
                        min_row=2, values_only=True)}
                    for name, *_, expected in cases:
                        with self.subTest(multi=multi, name=name):
                            if expected is None:
                                self.assertIsNone(actual[name])
                            else:
                                self.assertAlmostEqual(actual[name], expected)
                finally:
                    result.close()


if __name__ == "__main__":
    unittest.main()
