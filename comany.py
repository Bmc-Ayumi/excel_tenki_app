"""コマニー見積の加工と転記列対応。"""
import re
import openpyxl
import mojimoji
from openpyxl.styles import Alignment
from pyexcel import save_book_as



TRANSFER_COL_MAP = {'B': 'A', 'D': 'C', 'E': 'D', 'F': 'E', 'K': 'F', 'I': 'H'}


def comany(xlsx_path):
    # もし拡張子が .xls なら .xlsx に変換
    if xlsx_path.lower().endswith(".xls"):
        xlsx_converted = xlsx_path + "x"
        save_book_as(file_name=xlsx_path, dest_file_name=xlsx_converted)
        xlsx_path = xlsx_converted

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "内訳書" not in wb.sheetnames:
        wb.close()
        return None
    ws = wb["内訳書"]
    max_row = ws.max_row

    col = dict(b=2, c=3, f=6, g=7, i=9, j=10, k=11, l=12, m=13, n=14, o=15)

    # --- W/D/HをまとめてI列に出力 ---
    def fmt_dim(v):
        if v is None:
            return None
        try:
            f = float(v)
            return str(int(f)) if f.is_integer() else str(f)
        except:
            return str(v)

    def nonzero(v):
        if v in (None, "", 0, 0.0):
            return False
        try:
            return float(v) != 0
        except:
            return True

    for r in range(2, max_row + 1):
        w = ws.cell(r, col["i"]).value
        d = ws.cell(r, col["j"]).value
        h = ws.cell(r, col["k"]).value
        if w or d or h:
            merged = ""
            if nonzero(w): merged += f"W{fmt_dim(w)}"
            if nonzero(d): merged += f"D{fmt_dim(d)}"
            if nonzero(h): merged += f"H{fmt_dim(h)}"
            ws.cell(r, col["i"]).value = merged
            ws.merge_cells(start_row=r, start_column=col["i"], end_row=r, end_column=col["k"])
            ws.cell(r, col["i"]).alignment = Alignment(horizontal="left", vertical="center")

    # --- 【】内の文字をF列に抽出 ---
    pattern = re.compile(r"【(.*?)】")
    for r in range(2, max_row + 1):
        g_val = ws.cell(r, col["g"]).value
        g_str = str(g_val) if g_val else ""
        match = pattern.search(g_str)
        if match:
            ws.cell(r, col["f"]).value = match.group(1)

        if "施工" in g_str:
            l_val = ws.cell(r, col["l"]).value
            m_val = ws.cell(r, col["m"]).value
            try:
                merged = f"{float(l_val):.1f}{m_val}"
            except:
                merged = f"{l_val}{m_val}"
            ws.cell(r, col["i"]).value = merged
            ws.cell(r, col["l"]).value = 1
            ws.cell(r, col["m"]).value = "式"

    # --- ■や◆のある行をF列に転記 ---
    for r in range(2, max_row + 1):
        b_val = ws.cell(r, col["b"]).value
        c_val = ws.cell(r, col["c"]).value
        if b_val and "■" in str(b_val):
            ws.cell(r, col["f"]).value = mojimoji.han_to_zen(str(b_val))
        if c_val and "◆" in str(c_val):
            ws.cell(r, col["f"]).value = mojimoji.han_to_zen(str(c_val))

    # --- N列が空でO列が0以外 → NにOをコピー ---
    for r in range(2, max_row + 1):
        n_val = ws.cell(r, col["n"]).value
        o_val = ws.cell(r, col["o"]).value
        if (n_val is None or str(n_val).strip() == "") and nonzero(o_val):
            ws.cell(r, col["n"]).value = o_val

    # --- 列削除と挿入（元コードと同じ順序）---
    ws.delete_cols(10, 2)
    ws.delete_cols(7, 2)
    ws.delete_cols(1, 5)
    ws.insert_cols(2)

    # ★★★ A列が「部材」の行を特別処理 ★★★
    for r in range(2, ws.max_row + 1):
        a_val = ws.cell(r, 1).value  # A列（列削除後）
        if isinstance(a_val, str) and a_val.strip() == "部材":
            ws.cell(r, 1).value = "部材合計"
            for c in range(4, 8):  # D〜G
                ws.cell(r, c).value = None
    # ★★★ 完全な空行を削除（A〜H すべて空なら削除） ★★★
    for r in range(ws.max_row, 1, -1):  # 逆順で消す
        if all(ws.cell(r, c).value in (None, "") for c in range(1, 9)):  # A〜H列をチェック
            ws.delete_rows(r)


    # --- 不要シート削除（内訳書だけ残す）---
    for name in list(wb.sheetnames):
        if name != "内訳書":
            wb.remove(wb[name])

    output_path = xlsx_path.replace(".xlsx", "_comany_converted.xlsx")
    wb.save(output_path)
    wb.close()
    return output_path
